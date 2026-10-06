# -*- coding: utf-8 -*-
"""
Convertidor simple Markdown -> DOCX para los documentos de la defensa.
Uso: python generar_docx_alegaciones.py
Genera un .docx por cada .md de la lista, en la misma carpeta.
"""

import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ARCHIVOS = [
    "RESPUESTA_A_LA_EMPRESA_ALEGACIONES.md",
    "DENUNCIA_INSPECCION_TRABAJO.md",
    "DECLARACION_JANETH_SANCHEZ.md",
    "MI_DOCUMENTO_DE_PRUEBAS.md",
    "NOTAS_INTERNAS_VERIFICACION.md",
    "ALEGACIONES_JANETH_SANCHEZ.md",
    "ANEXO_I_INDICE_DE_PRUEBAS.md",
    "ANEXO_II_CUADRO_DE_COACCIONES.md",
    "ANEXO_III_CAMBIOS_DE_GRUPO_Y_RESIDENTES.md",
    "GUIA_PRESENTACION_Y_DEFENSA.md",
    "RESPUESTA_DENUNCIA_CARGO_POR_CARGO.md",
]


def insertar_imagen(doc, ruta, caption, ancho_cm=16.0):
    if not os.path.isfile(ruta):
        print("AVISO: no se encuentra la imagen", ruta)
        return
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    run.add_picture(ruta, width=Cm(ancho_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs_con_negritas(cap, caption)
        for r in cap.runs:
            r.italic = True
            r.font.size = Pt(9)


def add_runs_con_negritas(par, texto):
    partes = re.split(r"(\*\*.+?\*\*)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = par.add_run(parte[2:-2])
            run.bold = True
        else:
            par.add_run(parte)


def es_tabla_separadora(linea):
    return bool(re.match(r"^\s*\|[-:\s|]+\|\s*$", linea))


def parsear_tabla(filas):
    """filas: lista de líneas de tabla markdown (sin la separadora)."""
    datos = []
    for fila in filas:
        celdas = [c.strip() for c in fila.strip().strip("|").split("|")]
        datos.append(celdas)
    return datos


def convertir(md_path, docx_path):
    base_dir = os.path.dirname(os.path.abspath(md_path))
    with open(md_path, "r", encoding="utf-8") as f:
        lineas = f.read().split("\n")

    doc = Document()

    # Márgenes normales
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)

    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()

        # Imágenes incrustadas: ![título](ruta)
        m_img = re.match(r"^\s*!\[(.*?)\]\((.+?)\)\s*$", linea)
        if m_img:
            caption = m_img.group(1)
            ruta_img = m_img.group(2).strip().strip('"')
            if not os.path.isabs(ruta_img):
                ruta_img = os.path.join(base_dir, ruta_img)
            insertar_imagen(doc, ruta_img, caption)
            i += 1
            continue

        # Tablas
        if linea.strip().startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                if not es_tabla_separadora(lineas[i]):
                    bloque.append(lineas[i])
                i += 1
            datos = parsear_tabla(bloque)
            if datos:
                ncols = max(len(f) for f in datos)
                tabla = doc.add_table(rows=0, cols=ncols)
                try:
                    tabla.style = "Table Grid"
                except Exception:
                    pass
                for fila in datos:
                    celdas = tabla.add_row().cells
                    for j, texto in enumerate(fila):
                        if j >= ncols:
                            break
                        par = celdas[j].paragraphs[0]
                        par.text = ""
                        add_runs_con_negritas(par, texto)
                        for run in par.runs:
                            run.font.size = Pt(9)
                    if fila is datos[0]:
                        for celda in celdas:
                            for par in celda.paragraphs:
                                for run in par.runs:
                                    run.bold = True
            continue

        # Encabezados
        if linea.startswith("# "):
            h = doc.add_heading("", level=0)
            add_runs_con_negritas(h, linea[2:].strip())
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif linea.startswith("## "):
            h = doc.add_heading("", level=1)
            add_runs_con_negritas(h, linea[3:].strip())
        elif linea.startswith("### "):
            h = doc.add_heading("", level=2)
            add_runs_con_negritas(h, linea[4:].strip())
        # Separadores
        elif re.match(r"^-{3,}\s*$", linea):
            pass
        # Citas / notas
        elif linea.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            add_runs_con_negritas(par, linea.lstrip("> ").strip())
            for run in par.runs:
                run.italic = True
        # Listas con checkbox
        elif re.match(r"^\s*- \[[ xX]\]", linea):
            par = doc.add_paragraph(style="List Bullet")
            add_runs_con_negritas(par, linea)
        # Listas
        elif re.match(r"^\s*- ", linea):
            par = doc.add_paragraph(style="List Bullet")
            add_runs_con_negritas(par, linea.strip()[2:].strip())
        elif re.match(r"^\s*\d+\.\s", linea):
            par = doc.add_paragraph(style="List Number")
            add_runs_con_negritas(par, re.sub(r"^\s*\d+\.\s", "", linea))
        # Línea en blanco
        elif not linea.strip():
            pass
        # Párrafo normal
        else:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs_con_negritas(par, linea)
        i += 1

    doc.save(docx_path)


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    for archivo in ARCHIVOS:
        md = os.path.join(base, archivo)
        if not os.path.isfile(md):
            print("No encontrado:", md)
            continue
        docx = os.path.splitext(md)[0] + ".docx"
        convertir(md, docx)
        print("Generado:", os.path.basename(docx))


if __name__ == "__main__":
    main()
